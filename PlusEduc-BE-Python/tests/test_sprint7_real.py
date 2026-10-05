"""Sprint 7: dashboard, relatórios, filtros, pesquisa, exportação e histórico de operações (MongoDB real)."""
import io
import unicodedata
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pymongo import MongoClient

from app.core.auth import UserPrincipal
from app.core.config import Settings
from app.dependencies import get_current_user
from app.main import create_app

TEACHER = UserPrincipal(user_id="sprint7-teacher", email="teacher-sprint7@local", role="TEACHER")
STUDENT_ROLE = UserPrincipal(user_id="sprint7-student", email="student-sprint7@local", role="STUDENT", student_id="x")


def make_settings(**extra):
    return Settings(
        app_env="test-sprint7",
        mongodb_uri="mongodb://localhost:27017",
        mongodb_database="escola_db",
        mongodb_required=True,
        mongodb_server_selection_timeout_ms=3000,
        mongodb_connect_timeout_ms=3000,
        jwt_secret="test-sprint7-secret",
        **extra,
    )


@pytest.fixture
def ctx():
    mongo = MongoClient("mongodb://localhost:27017", serverSelectionTimeoutMS=3000)
    try:
        mongo.admin.command("ping")
        db = mongo["escola_db"]
        student = db.students.find_one({"active": {"$ne": False}, "$or": [{"class_id": {"$nin": [None, ""]}}, {"classId": {"$nin": [None, ""]}}]})
        if not student:
            pytest.fail("Nenhum aluno ativo matriculado em escola_db")
        classroom_id = str(student.get("class_id") or student.get("classId"))
        app = create_app(make_settings())
        app.dependency_overrides[get_current_user] = lambda: TEACHER
        with TestClient(app) as client:
            yield client, db, str(student["_id"]), classroom_id, student
    finally:
        mongo.close()


def norm(text):
    return unicodedata.normalize("NFD", text).encode("ascii", "ignore").decode().lower()


def real_grades_for_active_students(db):
    active = {str(s["_id"]) for s in db.students.find({"active": True}, {"_id": 1})}
    return [g for g in db.grades.find({}) if str(g.get("studentId")) in active]


def test_dashboard_kpis_and_charts_match_the_database(ctx):
    client, db, *_ = ctx
    payload = client.get("/api/dashboard/overview").json()
    grades = real_grades_for_active_students(db)
    values = [float(g.get("grade", g.get("gradeValue")) or 0) for g in grades]
    assert payload["kpis"]["students"] == db.students.count_documents({"active": True})
    assert payload["kpis"]["gradeLaunches"] == len(values)
    assert payload["kpis"]["averageGrade"] == pytest.approx(sum(values) / len(values), abs=0.01)
    assert sum(item["count"] for item in payload["distribution"]) == len(values)
    assert sum(item["launches"] for item in payload["averageByClassroom"]) == len(values)
    assert sum(item["launches"] for item in payload["averageBySubject"]) == len(values)
    assert sum(item["launches"] for item in payload["monthlyAverage"]) == len([g for g in grades if g.get("date")])
    assert payload["participation"]["delivered"] <= db.activity_submissions.count_documents({})
    assert len(payload["recentOperations"]) <= 8


def test_dashboard_filters_by_classroom_subject_and_period(ctx):
    client, db, student_id, classroom_id, _ = ctx
    scoped = client.get("/api/dashboard/overview", params={"classroomId": classroom_id}).json()
    assert [item["id"] for item in scoped["averageByClassroom"]] == [classroom_id]
    assert scoped["filters"]["Turma"]

    marker = f"Disciplina Sprint7 {uuid4().hex[:6]}"
    created = client.post("/api/grades", json={"studentId": student_id, "classroomId": classroom_id, "grade": 3.5, "attendance": False, "subject": marker, "activityType": "EXAM", "date": "2026-03-15T12:00:00Z"})
    assert created.status_code == 201
    try:
        by_subject = client.get("/api/dashboard/overview", params={"subject": marker.lower()}).json()
        assert by_subject["kpis"]["gradeLaunches"] == 1 and by_subject["kpis"]["averageGrade"] == 3.5
        assert by_subject["kpis"]["attendanceRate"] == 0
        assert by_subject["distribution"][1]["count"] == 1  # faixa 2 a 4
        assert by_subject["monthlyAverage"] == [{"month": "2026-03", "average": 3.5, "launches": 1}]
        window = client.get("/api/dashboard/overview", params={"subject": marker, "from": "2026-03-01", "to": "2026-03-31"}).json()
        assert window["kpis"]["gradeLaunches"] == 1
        outside = client.get("/api/dashboard/overview", params={"subject": marker, "from": "2026-04-01", "to": "2026-04-30"}).json()
        assert outside["kpis"]["gradeLaunches"] == 0 and outside["kpis"]["averageGrade"] is None
        assert client.get("/api/dashboard/overview", params={"from": "2026-05-01", "to": "2026-04-01"}).status_code == 422
        assert client.get("/api/dashboard/overview", params={"classroomId": "turma-inexistente"}).status_code == 404
    finally:
        assert client.delete(f"/api/grades/{created.json()['id']}").status_code == 204


def test_boletim_report_matches_the_grades_module(ctx):
    client, _, student_id, classroom_id, _ = ctx
    report = client.get("/api/reports/boletim", params={"classroomId": classroom_id}).json()
    summary = client.get(f"/api/grades/classroom/{classroom_id}/summary").json()
    expected = {row["studentId"]: row for row in summary["students"]}
    assert report["total"] == len(expected)
    by_name = {row["studentName"]: row for row in expected.values()}
    for row in report["rows"]:
        source = by_name[row["student"]]
        assert row["launches"] == source["launches"]
        assert (row["average"] or 0) == pytest.approx(source["average"], abs=0.01)
        assert row["attendanceRate"] == (source["attendanceRate"] if source["totalClasses"] else None)
    only_attention = client.get("/api/reports/boletim", params={"classroomId": classroom_id, "status": "ATENCAO"}).json()
    assert all(row["status"] == "Atenção" for row in only_attention["rows"])
    assert only_attention["total"] <= report["total"]


def test_reports_catalog_filters_and_errors(ctx):
    client, *_ = ctx
    catalog = client.get("/api/reports").json()
    assert {item["type"] for item in catalog} == {"boletim", "desempenho-disciplina", "alunos-atencao", "entregas-atividades", "historico-operacoes"}
    for item in catalog:
        assert client.get(f"/api/reports/{item['type']}").status_code == 200
    assert client.get("/api/reports/inexistente").status_code == 404
    deliveries = client.get("/api/reports/entregas-atividades", params={"status": "PENDENTE"}).json()
    assert all(row["deliveryRate"] is None or row["deliveryRate"] < 100 for row in deliveries["rows"])


def test_exports_in_csv_xlsx_and_pdf_contain_the_report_rows(ctx):
    client, _, _, classroom_id, _ = ctx
    params = {"classroomId": classroom_id}
    report = client.get("/api/reports/boletim", params=params).json()

    csv_response = client.get("/api/reports/boletim/export", params={**params, "format": "csv"})
    assert csv_response.status_code == 200 and "attachment" in csv_response.headers["content-disposition"]
    text = csv_response.content.decode("utf-8-sig").splitlines()
    header_index = next(i for i, line in enumerate(text) if line.startswith("Turma;Aluno;"))
    data_lines = [line for line in text[header_index + 1:] if line.strip() and ";" in line]
    assert len([line for line in data_lines if report["rows"][0]["classroom"] in line]) == report["total"]

    xlsx_response = client.get("/api/reports/boletim/export", params={**params, "format": "xlsx"})
    sheet = load_workbook(io.BytesIO(xlsx_response.content)).active
    header_row = next(r for r in range(1, 20) if sheet.cell(row=r, column=2).value == "Aluno")
    names = [sheet.cell(row=header_row + 1 + i, column=2).value for i in range(report["total"])]
    assert names == [row["student"] for row in report["rows"]]
    assert isinstance(sheet.cell(row=header_row + 1, column=4).value, (int, float))  # número de verdade, não texto

    pdf_response = client.get("/api/reports/boletim/export", params={**params, "format": "pdf"})
    assert pdf_response.content.startswith(b"%PDF") and pdf_response.headers["content-type"] == "application/pdf"

    assert client.get("/api/reports/boletim/export", params={"format": "docx"}).status_code == 422


def test_global_search_ignores_accents_and_case_and_is_staff_only(ctx):
    client, db, _, _, student = ctx
    name = student["name"]
    fragment = norm(name.split()[0])[:5].upper()
    found = client.get("/api/search", params={"q": fragment, "types": "students"}).json()
    assert any(hit["id"] == str(student["_id"]) for hit in found["results"])
    assert all(hit["type"] == "students" for hit in found["results"])
    assert client.get("/api/search", params={"q": student["matricula"]}).json()["results"][0]["id"] == str(student["_id"])
    classroom_name = db.classrooms.find_one({})["name"]
    assert any(hit["type"] == "classrooms" for hit in client.get("/api/search", params={"q": classroom_name[:8]}).json()["results"])
    assert client.get("/api/search", params={"q": "a"}).status_code == 422
    assert client.get("/api/search", params={"q": "zzzqqqxxx"}).json()["total"] == 0
    client.app.dependency_overrides[get_current_user] = lambda: STUDENT_ROLE
    for path in ("/api/search?q=maria", "/api/reports", "/api/reports/boletim", "/api/reports/boletim/export?format=csv", "/api/dashboard/overview", "/api/audit-logs"):
        assert client.get(path).status_code == 403, path


def test_audit_log_records_grade_changes_with_before_and_after_and_scopes_visibility():
    mongo = MongoClient("mongodb://localhost:27017", serverSelectionTimeoutMS=3000)
    db = mongo["escola_db"]
    token = uuid4().hex[:10]
    emails = [f"audit.a.{token}@professor.local", f"audit.b.{token}@professor.local", f"audit.s.{token}@aluno.local"]
    password = "Audit@12345"
    grade_ids = []
    try:
        student = db.students.find_one({"active": {"$ne": False}, "$or": [{"class_id": {"$nin": [None, ""]}}, {"classId": {"$nin": [None, ""]}}]})
        student_id, classroom_id = str(student["_id"]), str(student.get("class_id") or student.get("classId"))
        with TestClient(create_app(make_settings(audit_enabled=True))) as client:
            tokens = {}
            for email, kind in zip(emails, ("teacher", "teacher", "student")):
                response = client.post(f"/api/auth/register/{kind}", json={"name": f"Audit {kind} {token}", "email": email, "password": password})
                assert response.status_code == 200, response.text
            for email in emails:
                tokens[email] = client.post("/api/auth/login", json={"email": email, "password": password}).json()["accessToken"]
            a, b, s = ({"Authorization": f"Bearer {tokens[e]}"} for e in emails)

            created = client.post("/api/grades", headers=a, json={"studentId": student_id, "classroomId": classroom_id, "grade": 8.5, "attendance": True, "subject": "Auditoria", "activityType": "EXAM", "date": "2026-06-01T12:00:00Z"})
            assert created.status_code == 201
            grade_id = created.json()["id"]
            grade_ids.append(grade_id)
            assert client.put(f"/api/grades/{grade_id}", headers=a, json={"grade": 6.5}).status_code == 200
            assert client.delete(f"/api/grades/{grade_id}", headers=a).status_code == 204
            assert client.post("/api/auth/login", json={"email": emails[0], "password": "senha-errada"}).status_code == 401
            denied = client.post("/api/grades", headers=s, json={"studentId": student_id, "classroomId": classroom_id, "grade": 10, "date": "2026-06-01T12:00:00Z"})
            assert denied.status_code == 403

            page = client.get("/api/audit-logs", headers=a, params={"size": 50}).json()
            by_action = {}
            for item in page["items"]:
                by_action.setdefault((item["entity"], item["action"]), []).append(item)
            assert {item["actorEmail"] for item in page["items"]} == {emails[0]}  # só as próprias operações
            create_entry = by_action[("grade", "CREATE")][0]
            assert create_entry["entityId"] == grade_id and create_entry["success"] and "8,5" in create_entry["summary"]
            update_entry = by_action[("grade", "UPDATE")][0]
            assert update_entry["details"]["before"]["grade"] == 8.5 and update_entry["details"]["after"]["grade"] == 6.5
            assert "de 8,5 para 6,5" in update_entry["summary"]
            delete_entry = by_action[("grade", "DELETE")][0]
            assert delete_entry["details"]["before"]["grade"] == 6.5
            assert ("auth", "LOGIN") in by_action and by_action[("auth", "LOGIN")][0]["summary"] == f"Login de {emails[0]}"
            failed = by_action[("auth", "LOGIN_FAILED")][0]
            assert failed["success"] is False and failed["status"] == 401 and failed["actorEmail"] == emails[0]

            filtered = client.get("/api/audit-logs", headers=a, params={"entity": "grade", "action": "UPDATE"}).json()
            assert filtered["total"] == 1
            failures = client.get("/api/audit-logs", headers=a, params={"success": "false"}).json()
            assert failures["total"] >= 1 and all(not item["success"] for item in failures["items"])
            assert client.get("/api/audit-logs", headers=a, params={"q": "Auditoria"}).json()["total"] >= 1
            assert client.get("/api/audit-logs", headers=a, params={"from": "2026-01-01", "to": "2026-01-02"}).json()["total"] == 0

            # outro professor não vê o histórico do primeiro; aluno não acessa o histórico
            other = client.get("/api/audit-logs", headers=b, params={"entity": "grade"}).json()
            assert other["total"] == 0
            assert client.get("/api/audit-logs", headers=s).status_code == 403
            # a tentativa negada do aluno ficou registrada (visível ao administrador)
            denied_entry = db.audit_logs.find_one({"actorEmail": emails[2], "action": "ACCESS_DENIED"})
            assert denied_entry and denied_entry["status"] == 403 and denied_entry["entity"] == "grade"
    finally:
        db.grades.delete_many({"_id": {"$in": []}})
        from bson import ObjectId
        db.grades.delete_many({"_id": {"$in": [ObjectId(g) for g in grade_ids]}})
        db.audit_logs.delete_many({"actorEmail": {"$in": emails}})
        db.users.delete_many({"email": {"$in": emails}})
        db.teachers.delete_many({"email": {"$in": emails}})
        db.students.delete_many({"email": {"$in": emails}})
        mongo.close()


def test_audit_log_is_off_in_test_environment_by_default():
    mongo = MongoClient("mongodb://localhost:27017", serverSelectionTimeoutMS=3000)
    db = mongo["escola_db"]
    marker = f"audit.off.{uuid4().hex[:8]}@professor.local"
    try:
        with TestClient(create_app(make_settings())) as client:
            assert client.post("/api/auth/register/teacher", json={"name": "Audit desligado", "email": marker, "password": "Audit@12345"}).status_code == 200
        assert db.audit_logs.count_documents({"actorEmail": marker}) == 0
    finally:
        db.users.delete_many({"email": marker})
        db.teachers.delete_many({"email": marker})
        mongo.close()


def test_students_cannot_manage_classrooms_students_or_teachers(ctx):
    client, _, student_id, classroom_id, _ = ctx
    client.app.dependency_overrides[get_current_user] = lambda: STUDENT_ROLE
    body = {"name": "Turma Invasora", "year": 2026, "gradeLevel": "1º Ano", "teacherId": "x", "subjects": []}
    assert client.post("/api/classrooms", json=body).status_code == 403
    assert client.put(f"/api/classrooms/{classroom_id}", json={"name": "Renomeada"}).status_code == 403
    assert client.delete(f"/api/classrooms/{classroom_id}").status_code == 403
    assert client.post(f"/api/classrooms/{classroom_id}/enroll/{student_id}").status_code == 403
    assert client.delete(f"/api/classrooms/{classroom_id}/unenroll/{student_id}").status_code == 403
    for path in ("/api/classrooms", f"/api/classrooms/{classroom_id}", "/api/students", f"/api/students/{student_id}", "/api/teachers"):
        assert client.get(path).status_code == 403, path
