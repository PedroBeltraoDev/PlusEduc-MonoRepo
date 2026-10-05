from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.schemas.report import ReportColumn, ReportResponse
from app.services.audit_service import BRT

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _text(column: ReportColumn, value) -> str:
    """Texto de exibição (pt-BR) usado no CSV e no PDF."""
    if value is None:
        return "-"
    if column.type == "percent":
        return f"{value}%"
    if column.type == "decimal" and isinstance(value, (int, float)):
        return f"{value:.2f}".replace(".", ",")
    if column.type == "number" and isinstance(value, float):
        return f"{value:.2f}".replace(".", ",")
    if column.type == "date" and isinstance(value, str) and len(value) >= 10:
        year, month, day = value[:10].split("-")
        return f"{day}/{month}/{year}"
    if column.type == "datetime" and isinstance(value, str) and len(value) >= 16:
        date_part, time_part = value[:16].split(" ")
        year, month, day = date_part.split("-")
        return f"{day}/{month}/{year} {time_part}"
    return str(value)


def _local_generated(report: ReportResponse) -> str:
    return report.generatedAt.astimezone(BRT).strftime("%d/%m/%Y %H:%M")


def to_csv(report: ReportResponse) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([report.title])
    writer.writerow(["Gerado em", _local_generated(report)])
    for label, value in report.filters.items():
        writer.writerow([label, value])
    writer.writerow([])
    writer.writerow([column.label for column in report.columns])
    for row in report.rows:
        writer.writerow([_text(column, row.get(column.key)) for column in report.columns])
    writer.writerow([])
    for item in report.summary:
        writer.writerow([item.label, item.value])
    return buffer.getvalue().encode("utf-8-sig")  # BOM: o Excel abre os acentos corretamente


def to_xlsx(report: ReportResponse) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = report.title[:31]
    sheet.append([report.title])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append(["Gerado em", _local_generated(report)])
    for label, value in report.filters.items():
        sheet.append([label, value])
    sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append([column.label for column in report.columns])
    for cell in sheet[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1B365D")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in report.rows:
        sheet.append([_xlsx_value(column, row.get(column.key)) for column in report.columns])
    for index, column in enumerate(report.columns, start=1):
        letter = get_column_letter(index)
        for cell in sheet[letter][header_row:]:
            if column.type == "decimal" and isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
            elif column.type == "percent" and isinstance(cell.value, (int, float)):
                cell.number_format = '0"%"'
        width = max([len(column.label)] + [len(str(_text(column, r.get(column.key)))) for r in report.rows[:200]]) + 2
        sheet.column_dimensions[letter].width = min(max(width, 10), 60)
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
    if report.rows:
        sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(report.columns))}{header_row + len(report.rows)}"
    sheet.append([])
    for item in report.summary:
        sheet.append([item.label, item.value])
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()


def _xlsx_value(column: ReportColumn, value):
    if value is None:
        return None
    if column.type in {"date", "datetime"}:
        return _text(column, value)
    return value


def to_pdf(report: ReportResponse) -> bytes:
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm, title=report.title)
    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=8, leading=10)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    story = [Paragraph(report.title, styles["Title"]), Paragraph(report.description, styles["Normal"]), Spacer(1, 6)]
    meta = [f"Gerado em {_local_generated(report)}"] + [f"{label}: {value}" for label, value in report.filters.items()]
    story.append(Paragraph(" &nbsp;|&nbsp; ".join(meta), cell))
    story.append(Spacer(1, 8))
    if report.summary:
        story.append(Paragraph(" &nbsp;|&nbsp; ".join(f"<b>{item.label}:</b> {item.value}" for item in report.summary), styles["Normal"]))
        story.append(Spacer(1, 8))
    data = [[Paragraph(c.label, head) for c in report.columns]]
    for row in report.rows:
        data.append([Paragraph(_escape(_text(c, row.get(c.key))), cell) for c in report.columns])
    if len(data) == 1:
        story.append(Paragraph("Nenhum registro para os filtros selecionados.", styles["Normal"]))
    else:
        available = landscape(A4)[0] - 3 * cm
        weights = [max(len(c.label), 8) if c.type != "text" else 22 for c in report.columns]
        widths = [available * w / sum(weights) for w in weights]
        table = Table(data, colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B365D")), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C7D6EA")), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for index in range(2, len(data), 2):
            style.append(("BACKGROUND", (0, index), (-1, index), colors.HexColor("#EAF1F8")))
        table.setStyle(TableStyle(style))
        story.append(table)
        if report.total > len(report.rows):
            story.append(Paragraph(f"Exibindo {len(report.rows)} de {report.total} registros.", cell))
    doc.build(story)
    return out.getvalue()


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


EXPORTERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}
